#!/usr/bin/env python3
"""Phase 1 SEO data analysis for mydailytool.com.
Reads the GSC xlsx (Queries + Pages) and the GA4 Pages CSV, both in this folder.
Prints results to stdout. Read-only; writes nothing.
"""
import csv, re, os, warnings
warnings.filterwarnings("ignore")
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "mydailytool.com-Performance-on-Search-2026-06-17.xlsx")
CSV = os.path.join(HERE, "Pages_and_screens_Page_path_and_screen_class.csv")

wb = openpyxl.load_workbook(XLSX, data_only=True)

def rows(sheet):
    ws = wb[sheet]
    hdr = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        out.append(dict(zip(hdr, r)))
    return out

pages = rows("Pages")
queries = rows("Queries")

LC, PC = "Last 28 days Clicks", "Previous 28 days Clicks"
LI, PI = "Last 28 days Impressions", "Previous 28 days Impressions"
LP, PP = "Last 28 days Position", "Previous 28 days Position"

def num(x):
    return x if isinstance(x, (int, float)) else 0.0

bar = "=" * 78
def h(t): print("\n" + bar + "\n" + t + "\n" + bar)

# ---------------------------------------------------------------- 1. SITE TOTALS
h("1. SITE TOTALS (summed across GSC Pages sheet; position = impression-weighted avg)")
def totals(rs, ck, ik, pk):
    clk = sum(num(r[ck]) for r in rs)
    imp = sum(num(r[ik]) for r in rs)
    wpos = sum(num(r[pk]) * num(r[ik]) for r in rs) / imp if imp else 0
    return clk, imp, wpos
lc, li, lp = totals(pages, LC, LI, LP)
pc, pi, pp = totals(pages, PC, PI, PP)
print(f"{'metric':<14}{'last 28d':>12}{'prev 28d':>12}{'delta':>12}")
print(f"{'clicks':<14}{lc:>12.0f}{pc:>12.0f}{lc-pc:>+12.0f}")
print(f"{'impressions':<14}{li:>12.0f}{pi:>12.0f}{li-pi:>+12.0f}")
print(f"{'avg position':<14}{lp:>12.2f}{pp:>12.2f}{lp-pp:>+12.2f}  (lower=better)")
print(f"{'site CTR':<14}{(lc/li if li else 0):>11.2%}{(pc/pi if pi else 0):>12.2%}")
print(f"pages in report: {len(pages)}   queries in report: {len(queries)}")

# ---------------------------------------------------------------- 2. TIERS
h("2. PAGE TIERS  (Winning=pos 1-7 & clicks>0 | Promising=pos 8-20 & impr>0 | Dormant=rest)")
def tier(r):
    pos, clk, imp = num(r[LP]), num(r[LC]), num(r[LI])
    if pos == 0 and imp == 0:
        return "Dormant"
    if 1 <= pos <= 7 and clk > 0:
        return "Winning"
    if 8 <= pos <= 20 and imp > 0:
        return "Promising"
    return "Dormant"

for name in ("Winning", "Promising", "Dormant"):
    grp = [r for r in pages if tier(r) == name]
    grp.sort(key=lambda r: -num(r[LI]))
    print(f"\n--- {name}: {len(grp)} pages ---")
    print(f"{'URL':<52}{'impr':>7}{'pos':>8}{'Δimpr':>8}")
    for r in grp:
        u = r["Top pages"].replace("https://mydailytool.com", "") or "/"
        print(f"{u:<52}{num(r[LI]):>7.0f}{num(r[LP]):>8.1f}{num(r[LI])-num(r[PI]):>+8.0f}")

# ---------------------------------------------------------------- 3. MOVERS
h("3. BIGGEST MOVERS")
gain = sorted(pages, key=lambda r: -(num(r[LI]) - num(r[PI])))[:10]
print("\nTop 10 IMPRESSION GAINERS (last - prev):")
print(f"{'URL':<52}{'last':>7}{'prev':>7}{'Δ':>8}")
for r in gain:
    u = r["Top pages"].replace("https://mydailytool.com", "") or "/"
    print(f"{u:<52}{num(r[LI]):>7.0f}{num(r[PI]):>7.0f}{num(r[LI])-num(r[PI]):>+8.0f}")

# position loser = had a previous position AND position got worse (number went up)
losers = [r for r in pages if num(r[PP]) > 0 and num(r[LP]) > 0]
losers.sort(key=lambda r: -(num(r[LP]) - num(r[PP])))
print("\nTop 10 POSITION LOSERS (position worsened; both periods ranked):")
print(f"{'URL':<52}{'prev':>7}{'last':>7}{'Δpos':>8}")
for r in losers[:10]:
    u = r["Top pages"].replace("https://mydailytool.com", "") or "/"
    print(f"{u:<52}{num(r[PP]):>7.1f}{num(r[LP]):>7.1f}{num(r[LP])-num(r[PP]):>+8.1f}")

# ---------------------------------------------------------------- 4. QUERY CLUSTERS
h("4. QUERY CLUSTERS  (grouped by topic; flag big impr but best pos > 50)")
# cluster key -> (regex, target tool URL)
CLUSTERS = [
    ("capacitor / capacitance", r"capacitor|capacitance", "/tools/capacitance-calculator/"),
    ("days between dates",      r"days between|date difference|how many days", "/tools/days-between-dates/"),
    ("date calculator",         r"date calculat|add days|subtract days", "/tools/date-calculator/"),
    ("iso 8601 / 24:00",        r"iso ?8601|24:00|hh:mm|midnight|18:00", "/blog/iso-8601-cheat-sheet/"),
    ("timezone converter",      r"time ?zone|utc|gmt", "/tools/timezone-converter/"),
    ("unix timestamp",          r"unix|timestamp|epoch", "/tools/unix-timestamp-converter/"),
    ("json formatter",          r"json", "/tools/json-formatter/"),
    ("token counter / llm",     r"token|gpt|llm|tiktoken|claude", "/tools/token-counter/"),
    ("regex",                   r"regex|regular expression", "/tools/regex-tester/"),
    ("ohms law / electrical",   r"ohm|voltage|resistor|watt|amp|kva", "/tools/ohms-law-calculator/"),
    ("inductor / inductance",   r"inductor|inductance|reactance", "/tools/inductance-calculator/"),
    ("jwt",                     r"jwt|json web token", "/tools/jwt-decoder/"),
    ("base64",                  r"base64", "/tools/base64-encoder-decoder/"),
    ("sql formatter",           r"\bsql\b", "/tools/sql-formatter/"),
    ("yaml",                    r"yaml|yml", "/tools/yaml-formatter/"),
    ("cron",                    r"cron|crontab", "/tools/cron-expression-parser/"),
]
def clas(q):
    ql = q.lower()
    for name, rx, url in CLUSTERS:
        if re.search(rx, ql):
            return name, url
    return "(unclustered)", ""

agg = {}
for r in queries:
    q = r["Top queries"]
    name, url = clas(q)
    a = agg.setdefault(name, {"url": url, "impr": 0.0, "clk": 0.0, "best": 999.0, "n": 0, "top": []})
    a["impr"] += num(r[LI]); a["clk"] += num(r[LC]); a["n"] += 1
    if num(r[LI]) > 0:
        a["best"] = min(a["best"], num(r[LP]))
    a["top"].append((q, num(r[LI]), num(r[LP])))

ordered = sorted(agg.items(), key=lambda kv: -kv[1]["impr"])
print(f"{'cluster':<26}{'queries':>8}{'impr':>8}{'clk':>6}{'bestPos':>9}  target URL")
for name, a in ordered:
    flag = "  <-- DEMAND GAP (pos>50)" if a["impr"] >= 20 and a["best"] > 50 else ""
    print(f"{name:<26}{a['n']:>8}{a['impr']:>8.0f}{a['clk']:>6.0f}{a['best']:>9.1f}  {a['url']}{flag}")

print("\nTop query per flagged demand-gap cluster:")
for name, a in ordered:
    if a["impr"] >= 20 and a["best"] > 50:
        a["top"].sort(key=lambda t: -t[1])
        ex = "; ".join(f'"{q}" ({i:.0f}impr@{p:.0f})' for q, i, p in a["top"][:3])
        print(f"  {name}: {ex}")

# ---------------------------------------------------------------- 5. GSC vs GA4
h("5. GSC vs GA4 GAP")
ga = {}
with open(CSV, newline="") as f:
    started = False
    for row in csv.reader(f):
        if not row:
            continue
        if row[0] == "Page path and screen class":
            started = True; continue
        if started and row[0].startswith("/"):
            try:
                ga[row[0].rstrip("/") or "/"] = {"views": int(row[1]), "users": int(row[2])}
            except (ValueError, IndexError):
                pass

gsc = {}
for r in pages:
    path = r["Top pages"].replace("https://mydailytool.com", "")
    gsc[path.rstrip("/") or "/"] = {"clk": num(r[LC]), "impr": num(r[LI])}

print("\nA) GA4 views but ~0 GSC clicks (traffic that is NOT organic search):")
print(f"{'path':<46}{'GA4 views':>10}{'GSC clk':>9}{'GSC impr':>10}")
rowsA = []
for path, g in ga.items():
    s = gsc.get(path, {"clk": 0, "impr": 0})
    if g["views"] >= 5 and s["clk"] == 0:
        rowsA.append((path, g["views"], s["clk"], s["impr"]))
for path, v, c, i in sorted(rowsA, key=lambda t: -t[1])[:25]:
    print(f"{path:<46}{v:>10}{c:>9.0f}{i:>10.0f}")

print("\nB) GSC impressions but NO GA4 views (seen in search, never visited):")
print(f"{'path':<46}{'GSC impr':>10}{'GSC clk':>9}")
rowsB = []
for path, s in gsc.items():
    if s["impr"] > 0 and path not in ga:
        rowsB.append((path, s["impr"], s["clk"]))
for path, i, c in sorted(rowsB, key=lambda t: -t[1])[:25]:
    print(f"{path:<46}{i:>10.0f}{c:>9.0f}")

print("\nDone.")
